from __future__ import annotations

import csv
import html
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from typing import Any, Iterable

EMAIL_COLUMNS = [
    "Structure",
    "Format",
    "Rates Curve Index",
    "Issuer / Spread",
    "Currency",
    "Basket Type",
    "Size",
    "Bloomberg Ticker 1",
    "Bloomberg Ticker 2",
    "Bloomberg Ticker 3",
    "Bloomberg Ticker 4",
    "Bloomberg Ticker 5",
    "Reoffer / Upfront (%)",
    "Strike Date",
    "Tenor (m)",
    "Autocall (Yes/No)",
    "Frequency",
    "First Observation in (m)",
    "Autocall Trigger Level (%)",
    "Autocall Step Down/Up (%)",
    "Coupon Type",
    "Memory",
    "Coupon Trigger Level (%)",
    "Coupon Step Down/Up (%)",
    "Coupon p.a. (%)",
    "Strike Level (%)",
    "Downside Leverage (%)",
    "Barrier Type",
    "Barrier Level",
]

ALIASES = {
    "direction": ["direction", "side"],
    "size": ["size", "notional", "amount"],
    "assets": ["assets", "underlyings", "basket", "underlying assets"],
    "basket_type": ["basket type", "basket_type", "payoff basket type"],
    "maturity_date": ["maturity date", "maturity", "redemption date"],
    "currency": ["currency", "ccy"],
    "reference_date": ["reference date", "strike date", "initial valuation date"],
    "ko_barrier": ["ko barrier", "autocall barrier", "autocall trigger"],
    "coupon_barrier": ["coupon barrier", "coupon trigger"],
    "ki_barrier": ["ki barrier", "barrier level", "knock-in barrier"],
    "ki_monitoring": ["ki monitoring type", "barrier type", "knock-in monitoring"],
    "put_type": ["put type", "downside type"],
    "put_strike": ["put strike", "strike level"],
    "put_leverage": ["put leverage", "downside leverage"],
    "memory": ["memory", "memory coupon"],
    "coupons": ["coupons per observation", "coupon per observation", "coupon"],
    "coupon_dates": ["coupon dates", "payment dates"],
    "ko_dates": ["ko dates", "autocall dates", "observation dates"],
    "funding_type": ["funding type", "format"],
    "swap_benchmark": ["swap benchmark", "rates curve index", "benchmark"],
    "swap_spread": ["swap spread", "issuer spread", "spread"],
    "issue_date": ["issue date", "settlement date"],
    "reoffer": ["reoffer / upfront (%)", "reoffer", "upfront", "upfront (%)", "bid"],
    "coupon_pa": ["coupon p.a. (%)", "coupon pa", "annual coupon", "coupon p.a."],
    "strike_forward": ["strike forward (bd)", "strike forward", "bd"],
}

# Source "Basket Type" vocabulary -> Agile vocabulary.
BASKET_TYPES = {
    "BASKET": "Equally Weighted Basket",
    "EQUALLY WEIGHTED": "Equally Weighted Basket",
    "EQUALLY WEIGHTED BASKET": "Equally Weighted Basket",
    "EW": "Equally Weighted Basket",
    "EW BASKET": "Equally Weighted Basket",
    "WORST OF": "Worst Of",
    "WORSTOF": "Worst Of",
    "WO": "Worst Of",
    "BEST OF": "Best Of",
    "BESTOF": "Best Of",
    "BO": "Best Of",
}

COUNTERPARTY_SUBJECT = {
    "MORGAN STANLEY": "MS",
    "MS": "MS",
    "BNPP": "BNP",
    "BNP": "BNP",
    "SOCIETE GENERALE": "SG",
    "SG": "SG",
    "CITI": "CITI",
    "CITIBANK": "CITI",
    "JPM": "JPM",
    "JPMORGAN": "JPM",
    "BARCLAYS": "BARCLAYS",
    "HSBC": "HSBC",
    "UBS": "UBS",
}

# Placeholder Agile expects in a level schedule for an observation date where the
# feature (autocall / coupon) cannot trigger, e.g. a delayed first autocall.
SCHEDULE_NONE = "None"
SCHEDULE_SEPARATOR = "/"
# Two dates are treated as the same observation if they fall within this window.
DATE_MATCH_TOLERANCE = timedelta(days=10)


@dataclass
class MappingResult:
    fields: dict[str, str]
    source: dict[str, str]
    unmapped_source_fields: dict[str, str] = field(default_factory=dict)
    review_notes: list[str] = field(default_factory=list)
    trade_id: str = ""
    subject: str = "@AGILE"


def _normalise_label(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[_\-/]+", " ", text)
    text = re.sub(r"[^a-z0-9%(). ]+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def _clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%b-%y")
    return str(value).strip()


def _decode_csv_text(raw: bytes) -> str:
    """Decode CSV exports without introducing Unicode replacement characters.

    Excel/Windows CSV exports are frequently UTF-8, UTF-16, or Windows-1252.
    Decoding UTF-8 with ``errors="replace"`` turns otherwise valid punctuation
    and accented characters into U+FFFD, which Outlook often renders as ``?``.
    """
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    # latin-1 is exhaustive, so this is only a defensive fallback.
    return raw.decode("utf-8-sig")


def _read_rows(path_or_bytes: str | Path | bytes, filename: str | None = None) -> list[list[Any]]:
    if isinstance(path_or_bytes, bytes):
        raw = path_or_bytes
        suffix = Path(filename or "input.csv").suffix.lower()
    else:
        path = Path(path_or_bytes)
        raw = path.read_bytes()
        suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
    else:
        text = _decode_csv_text(raw)
        rows = list(csv.reader(io.StringIO(text)))

    rows = [row for row in rows if any(_clean_value(value) for value in row)]
    if not rows:
        raise ValueError("The file contains no usable rows.")
    return rows


def load_trade_file(
    path_or_bytes: str | Path | bytes,
    filename: str | None = None,
) -> list[tuple[dict[str, str], str]]:
    """Load one or many trades from CSV/XLSX.

    Supported layouts:
    1. Single key/value trade: field names in column A, values in column B.
    2. Multi-trade matrix: field names in column A, one trade per following column.
    3. Standard table: headers in the first row, one trade per following row.
    """
    rows = _read_rows(path_or_bytes, filename)
    max_columns = max(len(row) for row in rows)
    first_row = rows[0]
    first_cell = _clean_value(first_row[0] if first_row else "")
    first_row_trade_ids = [
        _clean_value(first_row[index] if index < len(first_row) else "")
        for index in range(1, max_columns)
    ]
    labelled_rows = sum(1 for row in rows[1:] if row and _clean_value(row[0]))

    # Matrix layout: blank top-left cell, trade IDs across row 1, fields down column A.
    is_matrix = (
        max_columns > 2
        and not first_cell
        and sum(bool(value) for value in first_row_trade_ids) >= 2
        and labelled_rows >= max(2, (len(rows) - 1) // 2)
    )
    if is_matrix:
        trades: list[tuple[dict[str, str], str]] = []
        for column_index in range(1, max_columns):
            trade_id = _clean_value(first_row[column_index] if column_index < len(first_row) else "")
            source: dict[str, str] = {}
            for row in rows[1:]:
                key = _clean_value(row[0] if row else "")
                value = _clean_value(row[column_index] if column_index < len(row) else "")
                if key:
                    source[key] = value
            if source and (trade_id or any(source.values())):
                trades.append((source, trade_id or f"Trade {column_index}"))
        if not trades:
            raise ValueError("No trades were found in the multi-column file.")
        return trades

    # Single key/value layout.
    key_value_rows = sum(
        1 for row in rows if len(row) >= 2 and _clean_value(row[0])
    )
    is_key_value = key_value_rows >= max(2, len(rows) // 2)
    if is_key_value:
        source: dict[str, str] = {}
        trade_id = ""
        for row_index, row in enumerate(rows):
            key = _clean_value(row[0] if row else "")
            value = _clean_value(row[1] if len(row) > 1 else "")
            if row_index == 0 and not key and value:
                trade_id = value
                continue
            if key:
                source[key] = value
        return [(source, trade_id)]

    # Standard tabular layout: first row contains headers and each later row is a trade.
    headers = [_clean_value(value) for value in rows[0]]
    trades = []
    trade_id_headers = {"trade id", "trade_id", "id", "name", "reference id"}
    for row_number, row in enumerate(rows[1:], start=1):
        source = {
            header: _clean_value(row[index] if index < len(row) else "")
            for index, header in enumerate(headers)
            if header
        }
        if not any(source.values()):
            continue
        trade_id = ""
        for header, value in source.items():
            if _normalise_label(header) in trade_id_headers and value:
                trade_id = value
                break
        trades.append((source, trade_id or f"Trade {row_number}"))
    if not trades:
        raise ValueError("The file does not contain any trade rows.")
    return trades


def load_key_value_file(
    path_or_bytes: str | Path | bytes,
    filename: str | None = None,
) -> tuple[dict[str, str], str]:
    """Backward-compatible helper that returns the first trade only."""
    return load_trade_file(path_or_bytes, filename)[0]


def _canonical_values(source: dict[str, str]) -> tuple[dict[str, str], dict[str, str]]:
    normalized_source = {_normalise_label(key): value for key, value in source.items()}
    used: set[str] = set()
    canonical: dict[str, str] = {}
    for canonical_name, aliases in ALIASES.items():
        for alias in aliases:
            normalized_alias = _normalise_label(alias)
            if normalized_alias in normalized_source:
                canonical[canonical_name] = normalized_source[normalized_alias]
                used.add(normalized_alias)
                break
    unmapped = {
        key: value
        for key, value in source.items()
        if _normalise_label(key) not in used
    }
    return canonical, unmapped


def _split_values(value: str) -> list[str]:
    text = str(value or "").strip()
    if not text or text.upper() in {"NA", "N/A", "NONE", "NULL"}:
        return []
    text = text.strip("[]")
    return [part.strip().strip("'\"") for part in text.split(",") if part.strip().strip("'\"")]


def _parse_date(value: str) -> datetime | None:
    value = str(value or "").strip()
    for date_format in ("%d-%b-%y", "%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(value, date_format)
        except ValueError:
            continue
    return None


def _parse_dates(value: str) -> list[datetime]:
    return [date for date in (_parse_date(part) for part in _split_values(value)) if date]


def _first_date(value: str) -> datetime | None:
    values = _split_values(value)
    return _parse_date(values[0]) if values else _parse_date(value)


def _months_between(start: datetime | None, end: datetime | None) -> str:
    if not start or not end:
        return ""
    months = (end.year - start.year) * 12 + end.month - start.month
    if end.day < start.day:
        months -= 1
    return str(max(months, 0))


def _number_text(text: str) -> str:
    """Return a compact numeric string ('100.00%' -> '100'), or the input if not numeric."""
    text = str(text or "").strip().replace("%", "").replace(",", "")
    if not text or text.upper() in {"NA", "N/A", "NONE"}:
        return ""
    try:
        return f"{float(text):g}"
    except ValueError:
        return text


def _percentage_number(value: str) -> str:
    values = _split_values(value)
    return _number_text(values[0] if values else str(value or ""))


def _percentage_numbers(value: str) -> list[float]:
    numbers: list[float] = []
    for part in _split_values(value):
        try:
            numbers.append(float(part.replace("%", "").replace(",", "")))
        except ValueError:
            continue
    return numbers


def _yes_no(value: str) -> str:
    text = str(value or "").strip().upper()
    if text in {"TRUE", "YES", "Y", "1"}:
        return "Yes"
    if text in {"FALSE", "NO", "N", "0"}:
        return "No"
    return str(value or "").strip()


def _tickers(assets: str) -> list[str]:
    result = []
    for asset in _split_values(assets):
        ticker = re.sub(r"\s+(Index|Equity|Curncy|Comdty|Corp|Govt)$", "", asset, flags=re.I).strip()
        result.append(ticker)
    return result[:5]


def _basket_type(value: str, ticker_count: int) -> tuple[str, str]:
    """Map the source basket type into Agile vocabulary.

    Returns (agile_value, review_note). The note is empty when nothing needs attention.
    """
    text = str(value or "").strip()
    key = re.sub(r"[^A-Z ]", " ", text.upper())
    key = re.sub(r"\s+", " ", key).strip()
    if key in BASKET_TYPES:
        return BASKET_TYPES[key], ""
    if key in {"", "NA", "N A", "NONE", "SINGLE"}:
        if ticker_count > 1:
            return "", "Basket Type is missing although several underlyings were found. Confirm the basket type with Agile."
        return "", ""
    # Unknown vocabulary: pass it through unchanged so the user can see and edit it.
    return text, f"Basket Type '{text}' is not in the known Agile vocabulary. Review before sending."


def _gap_months(left: datetime, right: datetime) -> int:
    months = (right.year - left.year) * 12 + right.month - left.month
    if right.day < left.day - 5:
        months -= 1
    return max(months, 0)


def _frequency(dates_value: str) -> tuple[str, str]:
    dates = _parse_dates(dates_value)
    if len(dates) < 2:
        return "", ""
    gaps = [_gap_months(left, right) for left, right in zip(dates, dates[1:])]
    gap = Counter(gaps).most_common(1)[0][0]
    labels = {1: "Monthly", 3: "Quarterly", 6: "Semi Annually", 12: "Annually"}
    return labels.get(gap, f"Every {gap} months"), str(gap)


def _observation_grid(
    coupon_dates: list[datetime],
    ko_dates: list[datetime],
    strike_date: datetime | None,
    maturity: datetime | None,
    gap_months: int,
) -> list[datetime]:
    """Return the observation calendar used to lay out level schedules.

    Preference order:
    1. Coupon dates (the denser calendar, normally starting at the first period).
    2. A synthetic calendar from the strike date at the detected frequency.
    3. The autocall dates themselves.
    """
    if coupon_dates:
        return coupon_dates
    if strike_date and maturity and gap_months > 0:
        grid: list[datetime] = []
        year, month = strike_date.year, strike_date.month
        while True:
            month += gap_months
            year += (month - 1) // 12
            month = (month - 1) % 12 + 1
            day = min(strike_date.day, 28)
            point = datetime(year, month, day)
            if point > maturity + DATE_MATCH_TOLERANCE:
                break
            grid.append(point)
        if grid:
            return grid
    return ko_dates


def _level_schedule(grid: list[datetime], dates: list[datetime], levels: list[float]) -> list[str]:
    """Lay out one level per grid observation, using SCHEDULE_NONE where no date matches."""
    schedule: list[str] = []
    for point in grid:
        match = None
        for index, date in enumerate(dates):
            if abs((date - point).days) <= DATE_MATCH_TOLERANCE.days:
                match = index
                break
        if match is None or match >= len(levels):
            schedule.append(SCHEDULE_NONE)
        else:
            schedule.append(f"{levels[match]:g}")
    return schedule


def _compact_schedule(schedule: list[str]) -> list[str]:
    """Drop trailing repeats: Agile carries the last level forward to maturity,
    so ``None/100/100/100/95/95/95`` is sent as ``None/100/100/100/95``."""
    compact = list(schedule)
    while len(compact) > 1 and compact[-1] == compact[-2]:
        compact.pop()
    return compact


def _trigger_and_step(
    levels_value: str,
    dates_value: str,
    grid: list[datetime],
) -> tuple[str, str, str]:
    """Return (trigger_level, step_down_up, note) following the Agile convention.

    * Flat schedule with no delayed start: trigger level = the level, step = "".
    * Anything else (step down/up or a delayed first observation): trigger level
      is left blank and the per-observation schedule goes in the step column,
      e.g. ``None/100/100/100/95``. Trailing repeats are dropped because Agile
      carries the last level forward to maturity.
    """
    levels = _percentage_numbers(levels_value)
    if not levels:
        return "", "", ""

    dates = _parse_dates(dates_value)
    schedule = _level_schedule(grid, dates, levels) if grid and dates else [f"{level:g}" for level in levels]

    delayed = schedule and schedule[0] == SCHEDULE_NONE
    flat = len(set(levels)) == 1
    if flat and not delayed:
        return f"{levels[0]:g}", "", ""

    note = ""
    if delayed:
        skipped = sum(1 for item in schedule if item == SCHEDULE_NONE)
        note = (
            f"Delayed start: the first {skipped} observation date{'s carry' if skipped != 1 else ' carries'} "
            f"no autocall and {'are' if skipped != 1 else 'is'} sent as '{SCHEDULE_NONE}'."
        )
    return "", SCHEDULE_SEPARATOR.join(_compact_schedule(schedule)), note


def _coupon_pa(coupons_value: str, coupon_pa_value: str, gap_months: int) -> tuple[str, str]:
    """Annualise coupons per observation, unless a coupon p.a. is already given."""
    explicit = _percentage_number(coupon_pa_value)
    if explicit:
        return explicit, ""
    coupons = _percentage_numbers(coupons_value)
    if not coupons:
        return "", ""
    if len(set(coupons)) > 1:
        return "Variable", "Coupons per observation vary across dates; Coupon p.a. needs a manual decision."
    if gap_months <= 0:
        return "", "Coupon p.a. could not be annualised because the coupon frequency is unknown."
    annual = coupons[0] * 12 / gap_months
    return f"{round(annual, 6):g}", ""


def _step(value: str) -> str:
    numbers = []
    for part in _split_values(value):
        try:
            numbers.append(float(part.replace("%", "")))
        except ValueError:
            return "Variable"
    if len(numbers) < 2 or len(set(numbers)) == 1:
        return ""
    differences = [round(right - left, 8) for left, right in zip(numbers, numbers[1:])]
    if len(set(differences)) == 1:
        return f"{differences[0]:g}"
    return "Variable"


def _subject_from_direction(direction: str) -> str:
    upper = str(direction or "").upper()
    for key, code in COUNTERPARTY_SUBJECT.items():
        if key in upper:
            return f"@AGILE {code}"
    token = re.split(r"\s+", upper.strip())[0] if upper.strip() else ""
    return f"@AGILE {token}".strip()


def map_trade(
    source: dict[str, str],
    trade_id: str = "",
    manual_overrides: dict[str, str] | None = None,
) -> MappingResult:
    canonical, unmapped = _canonical_values(source)
    notes: list[str] = []

    tickers = _tickers(canonical.get("assets", ""))
    strike_date = _first_date(canonical.get("reference_date", ""))
    maturity = _parse_date(canonical.get("maturity_date", ""))

    ko_dates_value = canonical.get("ko_dates", "")
    coupon_dates_value = canonical.get("coupon_dates", "")
    ko_dates = _parse_dates(ko_dates_value)
    coupon_dates = _parse_dates(coupon_dates_value)
    first_ko = ko_dates[0] if ko_dates else None
    has_ko = bool(_split_values(ko_dates_value))

    frequency, gap_text = _frequency(ko_dates_value)
    coupon_frequency, coupon_gap_text = _frequency(coupon_dates_value)
    if not frequency:
        frequency, gap_text = coupon_frequency, coupon_gap_text
    gap_months = int(gap_text) if gap_text.isdigit() else 0
    coupon_gap_months = int(coupon_gap_text) if coupon_gap_text.isdigit() else gap_months

    grid = _observation_grid(coupon_dates, ko_dates, strike_date, maturity, gap_months)

    has_put = bool(
        canonical.get("put_type", "")
        and canonical.get("put_type", "").upper() not in {"NA", "N/A", "NONE"}
    )

    if has_ko and has_put:
        structure = "Autocallable & BRC"
    elif has_ko:
        structure = "Autocallable"
    elif has_put:
        structure = "BRC"
    else:
        structure = ""
        notes.append("Structure could not be classified automatically.")

    knock_in = canonical.get("ki_barrier", "")
    knock_in_missing = not _split_values(knock_in)
    barrier_type = "None" if knock_in_missing else (canonical.get("ki_monitoring", "") or "Review")
    barrier_level = "" if knock_in_missing else _percentage_number(knock_in)

    coupon_barrier = canonical.get("coupon_barrier", "")
    coupon_barrier_value = _percentage_number(coupon_barrier)
    if coupon_barrier_value == "0":
        coupon_type = "Guaranteed"
    elif _split_values(coupon_barrier):
        coupon_type = "Conditional"
    else:
        coupon_type = "Fixed"

    basket_type, basket_note = _basket_type(canonical.get("basket_type", ""), len(tickers))
    if basket_note:
        notes.append(basket_note)

    autocall_trigger, autocall_step, autocall_note = _trigger_and_step(
        canonical.get("ko_barrier", ""), ko_dates_value, grid
    )
    if autocall_note:
        notes.append(autocall_note)

    coupon_trigger, coupon_step, coupon_note = _trigger_and_step(
        coupon_barrier, coupon_dates_value, grid
    )
    if coupon_note:
        notes.append(coupon_note.replace("autocall", "coupon"))

    coupon_pa, coupon_pa_note = _coupon_pa(
        canonical.get("coupons", ""), canonical.get("coupon_pa", ""), coupon_gap_months
    )
    if coupon_pa_note:
        notes.append(coupon_pa_note)

    fields = {column: "" for column in EMAIL_COLUMNS}
    fields.update(
        {
            "Structure": structure,
            "Format": canonical.get("funding_type", ""),
            "Rates Curve Index": canonical.get("swap_benchmark", ""),
            "Issuer / Spread": _percentage_number(canonical.get("swap_spread", "")),
            "Currency": canonical.get("currency", ""),
            "Basket Type": basket_type,
            "Size": re.sub(r"[^0-9.-]", "", canonical.get("size", "")),
            "Reoffer / Upfront (%)": _percentage_number(canonical.get("reoffer", "")),
            "Strike Date": strike_date.strftime("%d/%m/%Y") if strike_date else "",
            "Tenor (m)": _months_between(strike_date, maturity),
            "Autocall (Yes/No)": "Yes" if has_ko else "No",
            "Frequency": frequency,
            "First Observation in (m)": _months_between(strike_date, first_ko),
            "Autocall Trigger Level (%)": autocall_trigger,
            "Autocall Step Down/Up (%)": autocall_step,
            "Coupon Type": coupon_type,
            "Memory": _yes_no(canonical.get("memory", "")),
            "Coupon Trigger Level (%)": coupon_trigger,
            "Coupon Step Down/Up (%)": coupon_step,
            "Coupon p.a. (%)": coupon_pa,
            "Strike Level (%)": _percentage_number(canonical.get("put_strike", "")),
            "Downside Leverage (%)": _percentage_number(canonical.get("put_leverage", "")),
            "Barrier Type": barrier_type,
            "Barrier Level": barrier_level,
        }
    )
    for index, ticker in enumerate(tickers, start=1):
        fields[f"Bloomberg Ticker {index}"] = ticker

    if not fields["Strike Date"]:
        notes.append("Strike Date is missing. Agile will assume today's date; add one if the trade strikes on another day.")
    if not fields["Reoffer / Upfront (%)"]:
        notes.append("Reoffer / Upfront (%) is missing and needs a manual input or another source.")
    if has_put and not canonical.get("strike_forward"):
        notes.append("Strike Forward (Bd) is missing. Add it if Agile requires this field.")

    if manual_overrides:
        for key, value in manual_overrides.items():
            if key in fields:
                fields[key] = str(value)

    return MappingResult(
        fields=fields,
        source=source,
        unmapped_source_fields=unmapped,
        review_notes=notes,
        trade_id=trade_id,
        subject=_subject_from_direction(canonical.get("direction", "")),
    )


def map_trades(trades: Iterable[tuple[dict[str, str], str]]) -> list[MappingResult]:
    return [map_trade(source, trade_id) for source, trade_id in trades]


def _field_rows(fields_or_rows: dict[str, str] | Iterable[dict[str, str]]) -> list[dict[str, str]]:
    if isinstance(fields_or_rows, dict):
        return [fields_or_rows]
    return list(fields_or_rows)


def _html_cell_text(value: Any) -> str:
    text = "" if value is None else str(value)
    # Remove control characters that can corrupt HTML/MIME while preserving tabs/newlines.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    escaped = html.escape(text, quote=True)
    return escaped.replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>")


def build_html_table(fields_or_rows: dict[str, str] | Iterable[dict[str, str]]) -> str:
    """Build Outlook-safe HTML using inline styles only.

    Desktop Outlook uses the Microsoft Word rendering engine and ignores or
    inconsistently applies several modern CSS features. Keeping all critical
    table styling inline produces a much more stable result in Outlook and in
    downloaded .eml files.
    """
    rows = _field_rows(fields_or_rows)
    table_style = (
        "border-collapse:collapse;border-spacing:0;"
        "mso-table-lspace:0pt;mso-table-rspace:0pt;"
        "font-family:Arial,Helvetica,sans-serif;font-size:10px;color:#111111;"
    )
    header_style = (
        "border:1px solid #b7b7b7;background-color:#f2f2f2;"
        "padding:4px 6px;font-family:Arial,Helvetica,sans-serif;font-size:10px;"
        "font-weight:700;text-align:left;vertical-align:top;white-space:normal;"
    )
    cell_style = (
        "border:1px solid #b7b7b7;padding:4px 6px;"
        "font-family:Arial,Helvetica,sans-serif;font-size:10px;"
        "text-align:left;vertical-align:top;white-space:normal;"
    )
    headers = "".join(
        f'<th style="{header_style}">{_html_cell_text(column)}</th>'
        for column in EMAIL_COLUMNS
    )
    body_rows = []
    for fields in rows:
        values = "".join(
            f'<td style="{cell_style}">{_html_cell_text(fields.get(column, ""))}</td>'
            for column in EMAIL_COLUMNS
        )
        body_rows.append(f"<tr>{values}</tr>")
    body = "".join(body_rows)
    return f"""<!doctype html>
<html>
<head>
<meta http-equiv="Content-Type" content="text/html; charset=utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
</head>
<body style="margin:0;padding:0;font-family:Arial,Helvetica,sans-serif;font-size:10px;color:#111111;">
<table role="table" cellpadding="0" cellspacing="0" border="0" style="{table_style}">
<thead><tr>{headers}</tr></thead>
<tbody>{body}</tbody>
</table>
</body>
</html>"""


def build_eml(
    to_address: str,
    subject: str,
    fields_or_rows: dict[str, str] | Iterable[dict[str, str]],
) -> bytes:
    """Create an Outlook-friendly RFC 5322 message with explicit UTF-8 encoding."""
    message = EmailMessage(policy=SMTP)
    message["To"] = to_address
    message["Subject"] = subject
    message.set_content(
        "This message contains an HTML trade table. Please open it in an HTML-capable email client.",
        charset="utf-8",
        cte="base64",
    )
    message.add_alternative(
        build_html_table(fields_or_rows),
        subtype="html",
        charset="utf-8",
        cte="base64",
    )
    return message.as_bytes(policy=SMTP)


def combined_subject(results: Iterable[MappingResult]) -> str:
    subjects = {result.subject for result in results if result.subject}
    if len(subjects) == 1:
        return next(iter(subjects))
    if not subjects:
        return "@AGILE"
    return "@AGILE MULTI"


def save_outputs(
    results: MappingResult | Iterable[MappingResult],
    output_dir: str | Path,
    to_address: str = "Agile@marexfp.com",
) -> dict[str, Path]:
    result_list = [results] if isinstance(results, MappingResult) else list(results)
    if not result_list:
        raise ValueError("There are no mapped trades to save.")
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    html_path = output / "generated_trade_email.html"
    eml_path = output / "generated_trade_email.eml"
    json_path = output / "mapping_result.json"
    subject = combined_subject(result_list)
    field_rows = [result.fields for result in result_list]
    html_path.write_text(build_html_table(field_rows), encoding="utf-8")
    eml_path.write_bytes(build_eml(to_address, subject, field_rows))
    json_path.write_text(
        json.dumps(
            {
                "subject": subject,
                "trades": [
                    {
                        "trade_id": result.trade_id,
                        "fields": result.fields,
                        "review_notes": result.review_notes,
                        "unmapped_source_fields": result.unmapped_source_fields,
                    }
                    for result in result_list
                ],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return {"html": html_path, "eml": eml_path, "json": json_path}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Convert one or many trade rows into an Agile-style email.")
    parser.add_argument("input_file")
    parser.add_argument("--output-dir", default="output")
    parser.add_argument("--to", default="Agile@marexfp.com")
    args = parser.parse_args()

    loaded_trades = load_trade_file(args.input_file)
    mapped_results = map_trades(loaded_trades)
    outputs = save_outputs(mapped_results, args.output_dir, args.to)
    print(json.dumps({key: str(value) for key, value in outputs.items()}, indent=2))
    for result in mapped_results:
        if result.review_notes:
            print(f"\nReview notes for {result.trade_id or 'trade'}:")
            for note in result.review_notes:
                print(f"- {note}")
